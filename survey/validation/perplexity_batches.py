"""Split the blind coding sheet into text batches for Perplexity, and import its answers back.

    python3 perplexity_batches.py make                  # writes LLM_valid/batches/batch_NN.txt
    python3 perplexity_batches.py import "<model name>" # reads LLM_valid/perplexity/batch_NN_coded.txt into the sheet
"""
import re
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent / "src"))
from io_utils import ROOT, CODING, load_coder_sheet  # noqa: E402
from coding_schema import CODE_FIELDS  # noqa: E402

PPX = ROOT / "LLM_valid"
BATCH_SIZE = 50
RETURN_COLS = ["review_id"] + CODE_FIELDS + ["evidence_quote", "notes"]


def make():
    sheet = load_coder_sheet()
    # the 75 priority records first (both arms), then the rest; within that, sheet order
    sheet = pd.concat([sheet[sheet.priority_10pct == "1"], sheet[sheet.priority_10pct != "1"]])
    out = PPX / "batches"
    out.mkdir(parents=True, exist_ok=True)
    for i in range(0, len(sheet), BATCH_SIZE):
        chunk = sheet.iloc[i:i + BATCH_SIZE]
        n = i // BATCH_SIZE + 1
        lines = [f"BATCH {n:02d}: {len(chunk)} records ({chunk.review_id.iloc[0]} ... {chunk.review_id.iloc[-1]})", ""]
        for r in chunk.itertuples():
            lines += [f"### {r.review_id} | arm {r.arm} | {r.year}", f"Title: {r.title}", f"Abstract: {r.abstract}", ""]
        (out / f"batch_{n:02d}.txt").write_text("\n".join(lines))
    print(f"{(len(sheet) - 1) // BATCH_SIZE + 1} batches in {out}")


def read_returned(path: Path) -> pd.DataFrame:
    text = path.read_text().replace("```tsv", "").replace("```", "")
    rows = [l.split("\t") for l in text.splitlines() if l.strip() and not l.startswith("review_id")]
    bad = [r[0] for r in rows if len(r) != len(RETURN_COLS)]
    if bad:  # truncated or malformed rows are skipped (and reported as not returned), never guessed
        print(f"{path.name}: skipped {bad} (need {len(RETURN_COLS)} tab-separated fields)")
        rows = [r for r in rows if len(r) == len(RETURN_COLS)]
    df = pd.DataFrame(rows, columns=RETURN_COLS).apply(lambda c: c.str.strip())
    return df.replace("-", "")


def import_(model: str):
    files = sorted((PPX / "perplexity").glob("batch_*_coded.txt"))
    if not files:
        raise SystemExit("no batch_NN_coded.txt in LLM_valid/perplexity/")
    ret = pd.concat([read_returned(f).assign(_file=f.name) for f in files], ignore_index=True)
    sheet = load_coder_sheet()
    unknown = set(ret.review_id) - set(sheet.review_id)
    dup = ret[ret.review_id.duplicated(keep=False)]
    if unknown or len(dup):
        raise SystemExit(f"unknown ids {sorted(unknown)[:10]}; duplicated ids {sorted(set(dup.review_id))[:10]}")
    for n in sorted({int(re.search(r"batch_(\d+)", f.name).group(1)) for f in files}):
        want = [l[4:].split(" |")[0] for l in (PPX / "batches" / f"batch_{n:02d}.txt").read_text().splitlines()
                if l.startswith("### ")]
        missing = set(want) - set(ret.review_id)
        if missing:
            print(f"batch {n:02d}: {len(missing)} records not returned: {sorted(missing)}")

    wb = load_workbook(CODING / "coding_sheet.xlsx")
    ws = wb["coding"]
    col = {c.value: c.column for c in ws[1]}
    row_of = {ws.cell(r, col["review_id"]).value: r for r in range(2, ws.max_row + 1)}
    today = date.today().isoformat()
    for r in ret.itertuples(index=False):
        i = row_of[r.review_id]
        for f in RETURN_COLS[1:]:
            ws.cell(i, col[f]).value = getattr(r, f) or None
        ws.cell(i, col["coder_id"]).value = f"Perplexity: {model}"
        ws.cell(i, col["coding_date"]).value = today
    wb.save(CODING / "coding_sheet.xlsx")
    load_coder_sheet(CODING / "coding_sheet.xlsx").to_csv(CODING / "coding_sheet.tsv", sep="\t", index=False)
    print(f"imported {len(ret)} records from {len(files)} files; now run notebooks/03_validate_coding.ipynb")


if __name__ == "__main__":
    {"make": lambda: make(), "import": lambda: import_(sys.argv[2])}[sys.argv[1]]()
